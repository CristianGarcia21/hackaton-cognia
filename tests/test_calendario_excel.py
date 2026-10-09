"""Tests del MCP de calendario (#18), del Excel (#17) y de la ficha visual de la cita. Sin red."""

import asyncio
import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from server import tarjetas
from server.tools import calendario as K
from server.tools import citas as C
from server.tools import excel as X


def correr(coro):
    return asyncio.run(coro)


def manana_a(hora: str) -> str:
    d = date.today() + timedelta(days=1)
    d += timedelta(days=1) if d.weekday() == 6 else timedelta()
    return f"{d.isoformat()}T{hora}"


SEDE = {"sede_codigo": "7600102870-01", "sede_nombre": "FUNDACION VALLE DEL LILI", "motivo": "cirugía"}
CONTACTO = {"documento": "1020304050", "telefono": "3001234567"}


@pytest.fixture
def citas(tmp_path):
    return C.HerramientasCitas(tmp_path / "citas.db")


def registrar(citas, paciente="María Gómez", hora="10:00", **extra):
    return correr(citas.registrar_solicitud_cita(paciente=paciente, fecha_hora=manana_a(hora), **SEDE,
                                                 **(CONTACTO | extra)))


class GoogleFalso:
    def __init__(self, falla=False):
        self.creados, self.falla = [], falla

    async def crear(self, titulo, inicio, fin, descripcion, lugar):
        if self.falla:
            raise RuntimeError("403")
        self.creados.append((titulo, inicio, fin, descripcion))
        return "g1", "https://calendar.google.com/event?eid=abc"


def test_crea_evento_local_idempotente_y_feed_ics(citas, tmp_path):
    registrar(citas)
    cal = K.HerramientasCalendario(citas, tmp_path / "cal.db", google=None)
    texto = correr(cal.crear_evento_cita(solicitud_id=1))
    assert texto.startswith("Evento creado en el calendario de Kognia") and "pendiente de confirmación" in texto
    assert correr(cal.crear_evento_cita(solicitud_id="#1")).startswith("El evento ya existía")
    assert len(cal.eventos()) == 1
    feed = K.feed_ics(cal.eventos())
    assert "BEGIN:VEVENT" in feed and "María G." in feed and "María Gómez" not in feed  # feed público


def test_con_google_crea_el_evento_real_con_el_nombre_completo(citas, tmp_path):
    registrar(citas)
    google = GoogleFalso()
    cal = K.HerramientasCalendario(citas, tmp_path / "cal.db", google=google)
    texto = correr(cal.crear_evento_cita(solicitud_id=1))
    assert "en Google Calendar" in texto and "https://calendar.google.com/event?eid=abc" in texto
    titulo, inicio, fin, descripcion = google.creados[0]
    assert "pendiente de confirmación" in titulo and "María Gómez" in descripcion and (fin - inicio).seconds == 1800


def test_si_google_falla_queda_el_evento_local(citas, tmp_path):
    registrar(citas)
    cal = K.HerramientasCalendario(citas, tmp_path / "cal.db", google=GoogleFalso(falla=True))
    assert correr(cal.crear_evento_cita(solicitud_id=1)).startswith("Evento creado en el calendario de Kognia")


@pytest.mark.parametrize("sid,mensaje", [(99, "no existe"), ("abc", "no es un número")])
def test_errores_del_calendario_son_texto(citas, tmp_path, sid, mensaje):
    cal = K.HerramientasCalendario(citas, tmp_path / "cal.db", google=None)
    assert mensaje in correr(cal.crear_evento_cita(solicitud_id=sid))


def test_sin_fecha_exacta_no_hay_evento(citas, tmp_path):
    correr(citas.registrar_solicitud_cita(paciente="Ana", fecha_preferida="el lunes", **SEDE, **CONTACTO))
    cal = K.HerramientasCalendario(citas, tmp_path / "cal.db", google=None)
    assert "no tiene día y hora" in correr(cal.crear_evento_cita(solicitud_id=1))


def test_excel_una_hoja_por_ips_y_datos_enmascarados(citas):
    registrar(citas, documento="1234567890", telefono="3001234567")
    correr(citas.registrar_solicitud_cita(paciente="Luis", fecha_hora=manana_a("09:00"), sede_codigo="500102104-01",
                                          sede_nombre="HOSPITAL PABLO TOBON URIBE", motivo="consulta", **CONTACTO))
    libro = load_workbook(io.BytesIO(X.construir_excel(citas.solicitudes())))
    assert libro.sheetnames[0] == "Resumen" and len(libro.sheetnames) == 3
    filas = list(libro["FUNDACION VALLE DEL LILI"].values)
    assert filas[0][:3] == ("#", "Paciente", "Documento") and filas[1][2] == "***7890" and filas[1][3] == "***4567"
    texto = correr(X.HerramientasExcel(citas).exportar_solicitudes_excel())
    assert "2 solicitudes en 2 hojas" in texto and "/api/citas/excel" in texto


def test_proponer_cita_valida_sin_registrar(citas):
    texto = correr(citas.proponer_cita(paciente="Ana", fecha_hora=manana_a("10:00"), **SEDE, **CONTACTO))
    assert texto.startswith("Propuesta lista") and citas.solicitudes() == []
    registrar(citas, paciente="Otro", hora="10:00")
    ocupada = correr(citas.proponer_cita(paciente="Ana", fecha_hora=manana_a("10:00"), **SEDE, **CONTACTO))
    assert ocupada.startswith("Error:") and "ya tiene una solicitud" in ocupada


def test_fichas_de_cada_fase():
    args = {"paciente": "María Gómez", "fecha_hora": manana_a("10:30"), **SEDE}
    p = tarjetas.tarjeta("proponer_cita", args, "Propuesta lista ...")
    assert p.kind == "cita" and p.data["fase"] == "por_confirmar" and p.data["hora"] == "10:30"
    assert p.data["paciente"] == "María G." and p.data["sede"] == "FUNDACION VALLE DEL LILI"
    r = tarjetas.tarjeta("registrar_solicitud_cita", args, "Solicitud #7 registrada ...")
    assert r.data["fase"] == "registrada" and r.data["solicitud_id"] == 7 and r.data["ics"] == "/api/citas/7.ics"
    e = tarjetas.tarjeta("crear_evento_cita", {"solicitud_id": 7},
                         "Evento creado en Google Calendar para la solicitud #7 ... Enlace: https://calendar.google.com/x")
    assert e.kind == "evento" and e.data["google"] == "https://calendar.google.com/x"
    assert tarjetas.tarjeta("proponer_cita", args, "Error: ocupada") is None
    assert tarjetas.tarjeta("buscar_ips", {}, "x") is None


def test_exportar_excel_deja_el_boton_de_descarga():
    a = tarjetas.tarjeta("exportar_solicitudes_excel", {}, "Excel listo con 4 solicitudes en 2 hojas por IPS. Se descarga")
    assert a.kind == "excel" and a.link == "/api/citas/excel" and a.data["filas"] == 4


def test_excel_de_una_sola_ips_listo_para_enviar(citas):
    registrar(citas)
    correr(citas.registrar_solicitud_cita(paciente="Luis", fecha_hora=manana_a("09:00"), sede_codigo="500102104-01",
                                          sede_nombre="HOSPITAL PABLO TOBON URIBE", motivo="consulta", **CONTACTO))
    libro = load_workbook(io.BytesIO(X.construir_excel(citas.solicitudes(), sede="7600102870-01")))
    assert libro.sheetnames == ["FUNDACION VALLE DEL LILI"]
    filas = list(libro.active.values)
    assert "FUNDACION VALLE DEL LILI" in filas[0][0] and "1 solicitud" in filas[1][0]
    assert filas[4][:2] == ("#", "Paciente") and filas[5][1] == "María Gómez" and len(filas) == 6  # solo esa IPS
    assert X.nombre_archivo(citas.solicitudes(), "7600102870-01") == "solicitudes-fundacion-valle-del-lili.xlsx"


def test_nombre_de_archivo_conserva_las_letras_con_tilde():
    filas = [{"sede_codigo": "1", "sede_nombre": "Fundación Clínica Ñuñoa"}]
    assert X.nombre_archivo(filas, "1") == "solicitudes-fundacion-clinica-nunoa.xlsx"


@pytest.mark.parametrize("sede", ['"', "漢", "\r\n"])
def test_nombre_de_archivo_nunca_usa_el_parametro_crudo(sede):
    assert X.nombre_archivo([], sede) == "solicitudes-ips.xlsx"
