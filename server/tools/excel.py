"""Tools de Excel (issue #17, spec §8): las solicitudes de cita organizadas para enviarlas a cada IPS.

    xl = HerramientasExcel(citas)
    await xl.exportar_solicitudes_excel()     # -> texto con el enlace /api/citas/excel
    construir_excel(citas.solicitudes())      # -> bytes del .xlsx (lo sirve GET /api/citas/excel)

Una hoja "Resumen" y una hoja por IPS, lista para enviar. La URL es pública: documento y teléfono van
enmascarados (últimos 4 dígitos) salvo que EXCEL_DATOS_COMPLETOS=1.
"""

import asyncio
import io
import os
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.tools import Tool, tool
from server.tools.citas import HerramientasCitas

ENCABEZADO = ["#", "Paciente", "Documento", "Teléfono", "Motivo", "Fecha y hora", "Preferencia dicha", "Estado",
              "Registrada"]
COLOR = "0E7490"  # --primario-fuerte de la UI


def _ocultar(v: str | None) -> str:
    if not v:
        return ""
    if os.getenv("EXCEL_DATOS_COMPLETOS") == "1":
        return v
    digitos = re.sub(r"\D", "", v)
    return f"***{digitos[-4:]}" if digitos else "***"


def _hoja_valida(nombre: str, usadas: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", nombre).strip()[:28] or "Sede"
    nombre, n = base, 2
    while nombre in usadas:
        nombre, n = f"{base[:25]} {n}", n + 1
    usadas.add(nombre)
    return nombre


def _formato(ws, anchos: list[int]) -> None:
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=COLOR)
        celda.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def nombre_archivo(filas: list[dict], sede: str | None) -> str:
    if not sede:
        return "solicitudes-por-ips.xlsx"
    nombre = next((f["sede_nombre"] for f in filas if f["sede_codigo"] == sede and f.get("sede_nombre")), sede)
    sin_tildes = unicodedata.normalize("NFKD", nombre).encode("ascii", "ignore").decode()  # "Fundación" → "Fundacion"
    limpio = re.sub(r"[^A-Za-z0-9]+", "-", sin_tildes).strip("-").lower()
    return f"solicitudes-{limpio or 'ips'}.xlsx"  # nunca el parámetro crudo: va en una cabecera HTTP


def construir_excel(filas: list[dict], sede: str | None = None) -> bytes:
    """Todas las IPS (Resumen + una hoja por sede) o, con `sede`, solo las solicitudes de esa IPS listas para
    enviarle: una hoja con encabezado (para quién es, cuántas, estado) y la tabla."""
    if sede:
        return _excel_de_una_sede([f for f in filas if f["sede_codigo"] == sede], sede)
    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen"
    resumen.append(["IPS (sede)", "Id sede", "Solicitudes", "Estado"])
    por_sede: dict[str, list[dict]] = {}
    for f in sorted(filas, key=lambda x: (x.get("fecha_hora") or "9999", x["id"])):
        por_sede.setdefault(f["sede_codigo"], []).append(f)
    usadas = {"Resumen"}
    for codigo, lista in sorted(por_sede.items(), key=lambda kv: -len(kv[1])):
        nombre = next((x["sede_nombre"] for x in lista if x.get("sede_nombre")), f"Sede {codigo}")
        resumen.append([nombre, codigo, len(lista), "listas para enviar a la IPS"])
        ws = libro.create_sheet(_hoja_valida(nombre, usadas))
        ws.append(ENCABEZADO)
        for x in lista:
            ws.append([x["id"], x["paciente"], _ocultar(x.get("documento")), _ocultar(x.get("telefono")), x["motivo"],
                       (x.get("fecha_hora") or "").replace("T", " "), x.get("fecha_preferida") or "", x["estado"],
                       (x.get("creada") or "").replace("T", " ")])
        _formato(ws, [6, 26, 14, 14, 32, 18, 22, 34, 20])
    _formato(resumen, [40, 18, 12, 30])
    if not por_sede:
        resumen.append(["Todavía no hay solicitudes", "", 0, ""])
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _excel_de_una_sede(filas: list[dict], sede: str) -> bytes:
    from datetime import datetime
    libro = Workbook()
    ws = libro.active
    nombre = next((f["sede_nombre"] for f in filas if f.get("sede_nombre")), f"Sede {sede}")
    ws.title = _hoja_valida(nombre, set())
    ws.append([f"Solicitudes de cita para: {nombre}"])
    ws.append([f"Id de sede (REPS): {sede}  ·  {len(filas)} solicitud{'es' if len(filas) != 1 else ''}  ·  "
               f"Estado: pendientes de confirmación por la IPS"])
    ws.append([f"Generado por Kognia el {datetime.now():%Y-%m-%d %H:%M} con datos de datos.gov.co"])
    ws.append([])
    ws.append(ENCABEZADO)
    for x in sorted(filas, key=lambda f: (f.get("fecha_hora") or "9999", f["id"])):
        ws.append([x["id"], x["paciente"], _ocultar(x.get("documento")), _ocultar(x.get("telefono")), x["motivo"],
                   (x.get("fecha_hora") or "").replace("T", " "), x.get("fecha_preferida") or "", x["estado"],
                   (x.get("creada") or "").replace("T", " ")])
    if not filas:
        ws.append(["", "No hay solicitudes para esta IPS"])
    for i, ancho in enumerate([6, 26, 14, 14, 32, 18, 22, 34, 20], 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws["A1"].font = Font(bold=True, size=14, color=COLOR)
    ws["A2"].font = Font(color="475569")
    ws["A3"].font = Font(italic=True, color="5B6779")
    for celda in ws[5]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=COLOR)
    ws.freeze_panes = "A6"
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


class HerramientasExcel:
    def __init__(self, citas: HerramientasCitas):
        self.citas = citas

    async def exportar_solicitudes_excel(self) -> str:
        """Genera el Excel con todas las solicitudes de cita organizadas por IPS (una hoja por sede), listo para enviarse a cada IPS. Úsala cuando el usuario pida exportar, descargar o enviar las solicitudes."""
        try:
            filas = await asyncio.to_thread(self.citas.solicitudes)
            await asyncio.to_thread(construir_excel, filas)  # valida que se pueda generar
        except Exception:  # noqa: BLE001
            return "Error: no pude generar el Excel en este momento. Intenta de nuevo."
        sedes = len({f["sede_codigo"] for f in filas})
        return (f"Excel listo con {len(filas)} solicitud{'es' if len(filas) != 1 else ''} en {sedes} hoja"
                f"{'s' if sedes != 1 else ''} por IPS, para enviar a cada una. Se descarga en /api/citas/excel; "
                "dile al usuario que el botón de descarga está en pantalla.")

    def tools(self) -> list[Tool]:
        return [tool(self.exportar_solicitudes_excel)]
